#!/usr/bin/env python3
"""Single entry point for the MS100 SDMT/MSFC/NODDI reproduction workflow.

Large and participant-level files are always written below ``--work-root``.
The package directory is treated as read-only distribution material; it is
never used as a work directory and the source BIDS tree is read-only. AMICO and ANTs create and
register the NODDI maps; MATLAB then generates every paper-facing metric table,
refits all 196 qMRI models, and supplies the figure inputs.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import numbers
import os
import shutil
import subprocess
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Sequence


PACKAGE_ROOT = Path(__file__).resolve().parents[1]
HELPER_ROOT = PACKAGE_ROOT / "helper"
PIPELINE_ROOT = HELPER_ROOT / "pipeline"
CONFIG_ROOT = HELPER_ROOT / "config"
ANALYSIS_CONFIG = CONFIG_ROOT / "analysis_config.json"
MATLAB_METRIC_TABLE_RUNNER = PACKAGE_ROOT / "code" / "generate_all_metric_tables.m"
MATLAB_RUNNER = PACKAGE_ROOT / "code" / "run_all_qmri_manuscript_models.m"
MATLAB_FIGURE_RUNNER = PACKAGE_ROOT / "code" / "figure6_r2_comparison_sdmt_msfc_noddi.m"
MATLAB_TABLE_RUNNER = PACKAGE_ROOT / "code" / "generate_manuscript_tables.m"
VOLUMES_ROOT = Path("/Volumes")
FIT_HASHED_OUTPUTS = (
    "fit_NDI.nii.gz",
    "fit_ODI.nii.gz",
    "fit_FWF.nii.gz",
    "fit_dir.nii.gz",
    "fit_RMSE.nii.gz",
    "fit_NRMSE.nii.gz",
    "config.pickle",
    "acquisition.scheme",
    "run.log",
    "subject_metrics.json",
    "qc_slice.npz",
)


class ReproductionError(RuntimeError):
    """A fail-closed validation or pipeline error."""


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(8 * 1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def atomic_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".partial")
    temporary.write_text(
        json.dumps(payload, indent=2, sort_keys=True, default=str) + "\n",
        encoding="utf-8",
    )
    os.replace(temporary, path)


def atomic_text(path: Path, payload: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".partial")
    temporary.write_text(payload, encoding="utf-8")
    os.replace(temporary, path)


def is_relative_to(path: Path, parent: Path) -> bool:
    try:
        path.relative_to(parent)
        return True
    except ValueError:
        return False


def resolve_existing(path: Path | None, label: str, directory: bool = False) -> Path:
    if path is None:
        raise ReproductionError(f"{label} is required")
    resolved = path.expanduser().resolve()
    exists = resolved.is_dir() if directory else resolved.is_file()
    if not exists:
        kind = "directory" if directory else "file"
        raise ReproductionError(f"Missing {label} {kind}: {resolved}")
    return resolved


def validate_work_root(work_root: Path, bids_root: Path | None) -> Path:
    root = work_root.expanduser().resolve()
    forbidden = {Path("/"), Path.home().resolve(), PACKAGE_ROOT.resolve()}
    if root in forbidden:
        raise ReproductionError(f"Unsafe --work-root: {root}")
    if root == VOLUMES_ROOT or VOLUMES_ROOT in root.parents:
        raise ReproductionError("--work-root must be local and cannot be below /Volumes")
    if is_relative_to(root, PACKAGE_ROOT.resolve()) or is_relative_to(PACKAGE_ROOT.resolve(), root):
        raise ReproductionError("--work-root must be outside the public package")
    if bids_root is not None:
        source = bids_root.resolve()
        if root == source or is_relative_to(root, source) or is_relative_to(source, root):
            raise ReproductionError("Source and work trees must be separate")
    return root


def package_hashes() -> dict[str, str]:
    paths: list[Path] = []
    for path in PACKAGE_ROOT.rglob("*"):
        if not path.is_file():
            continue
        relative = path.relative_to(PACKAGE_ROOT)
        if (
            "__pycache__" in relative.parts
            or path.name == ".DS_Store"
            or path.suffix in {".pyc", ".partial"}
            or path.name.endswith(".inspect.ndjson")
        ):
            continue
        paths.append(path)
    return {
        str(path.relative_to(PACKAGE_ROOT)): sha256_file(path)
        for path in sorted(paths)
    }


def run_command(
    command: Sequence[str],
    *,
    manifest: dict[str, Any],
    stage: str,
    env: dict[str, str] | None = None,
    cwd: Path | None = None,
    dry_run: bool = False,
) -> None:
    printable = " ".join(shlex_quote(part) for part in command)
    record: dict[str, Any] = {
        "stage": stage,
        "started_utc": utc_now(),
        "command": printable,
        "cwd": str(cwd or Path.cwd()),
    }
    manifest.setdefault("stages", []).append(record)
    print(f"\n[{stage}] {printable}", flush=True)
    if dry_run:
        record.update(status="DRY_RUN", finished_utc=utc_now())
        return
    completed = subprocess.run(
        list(command),
        cwd=str(cwd) if cwd else None,
        env=env,
        check=False,
    )
    record.update(
        status="PASS" if completed.returncode == 0 else "FAIL",
        return_code=completed.returncode,
        finished_utc=utc_now(),
    )
    if completed.returncode != 0:
        raise ReproductionError(f"Stage {stage} failed with exit code {completed.returncode}")


def shlex_quote(value: str) -> str:
    import shlex

    return shlex.quote(str(value))


def required_source_paths(bids_root: Path) -> list[Path]:
    return [
        bids_root / "derivatives" / "TractoFlow" / "ses-01",
        bids_root / "derivatives" / "TractoFlow_post",
        bids_root / "derivatives" / "freesurfer",
        bids_root / "derivatives" / "maps",
        bids_root / "derivatives" / "lesionMask",
        bids_root / "derivatives" / "derivativesStats" / "clinicalScore.xlsx",
        bids_root / "derivatives" / "derivativesStats" / "GroupTractT1_All.xlsx",
        bids_root / "derivatives" / "derivativesStats" / "GroupTractMTR_All.xlsx",
        bids_root / "derivatives" / "derivativesStats" / "GroupTractFA_All.xlsx",
        bids_root / "derivatives" / "derivativesStats" / "GroupTractMD_All.xlsx",
        bids_root / "derivatives" / "derivativesStats" / "icometrixLesionStats.xlsx",
    ]


def analysis_settings() -> dict[str, Any]:
    return json.loads(ANALYSIS_CONFIG.read_text(encoding="utf-8"))


def canonical_subject_payload(subjects: Sequence[str]) -> bytes:
    return ("\n".join(sorted(subjects)) + "\n").encode("utf-8")


def validate_source_cohort(
    bids_root: Path, subjects: Sequence[str], label: str
) -> dict[str, Any]:
    study = analysis_settings()["study"]
    observed_count = len(subjects)
    expected_count = int(study["expected_imaging_subjects"])
    manifest_lines: list[str] = []
    for subject in sorted(subjects):
        affine = (
            bids_root
            / "derivatives"
            / "TractoFlow"
            / "ses-01"
            / subject
            / "Register_T1"
            / f"{subject}__output0GenericAffine.mat"
        )
        if not affine.is_file():
            raise ReproductionError(
                f"{label} is missing a required registration transform; "
                "participant labels are intentionally not printed by the public runner."
            )
        manifest_lines.append(f"{subject}/{affine.name}\t{sha256_file(affine)}")
    observed_hash = hashlib.sha256(
        ("\n".join(manifest_lines) + "\n").encode("utf-8")
    ).hexdigest()
    expected_hash = str(study["expected_source_cohort_manifest_sha256"])
    if observed_count != expected_count or observed_hash != expected_hash:
        raise ReproductionError(
            f"{label} does not match the frozen content-bound cohort manifest: "
            f"expected count={expected_count} and SHA-256={expected_hash}; "
            f"observed count={observed_count} and SHA-256={observed_hash}. "
            "Participant labels are intentionally not printed by the public runner."
        )
    return {
        "subjects": observed_count,
        "source_cohort_manifest_sha256": observed_hash,
        "status": "PASS",
    }


def validate_source_tree(bids_root: Path) -> dict[str, Any]:
    missing = [str(path) for path in required_source_paths(bids_root) if not path.exists()]
    if missing:
        raise ReproductionError("Required private inputs are missing:\n  " + "\n  ".join(missing))
    source_root = bids_root / "derivatives" / "TractoFlow" / "ses-01"
    observed = sorted(path.name for path in source_root.glob("sub-*") if path.is_dir())
    report = validate_source_cohort(bids_root, observed, "TractoFlow source cohort")
    report["source_root"] = str(source_root)
    return report


def is_semantic_missing(value: Any) -> bool:
    if value is None:
        return True
    if isinstance(value, str):
        return not value.strip()
    return isinstance(value, numbers.Real) and math.isnan(float(value))


def semantic_legacy_tract_fingerprint(path: Path) -> dict[str, Any]:
    """Hash legacy tract cells at five significant figures.

    The source maps are stored as single precision, so an active MATLAB
    reaggregation can differ from the historical workbook by one float32
    step even when both are scientifically identical. Five significant
    figures remain finer than the manuscript display precision while the
    payload still binds every subject, column, and missing cell.
    """
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook[workbook.sheetnames[0]]
        row_iterator = worksheet.iter_rows(values_only=True)
        try:
            first_row = next(row_iterator)
        except StopIteration as failure:
            raise ReproductionError(f"Legacy tract workbook is empty: {path.name}") from failure
        headers = ["" if value is None else str(value).strip() for value in first_row]
        while headers and not headers[-1]:
            headers.pop()
        if len(headers) < 2 or headers[0] != "SubjectID" or any(not value for value in headers):
            raise ReproductionError(f"Invalid legacy tract header in {path.name}")
        if len(set(headers)) != len(headers):
            raise ReproductionError(f"Duplicate legacy tract headers in {path.name}")

        lines: list[str] = []
        subjects: set[str] = set()
        ignored_all_empty_rows = 0
        for raw_row in row_iterator:
            values = list(raw_row[: len(headers)])
            values.extend([None] * (len(headers) - len(values)))
            if all(is_semantic_missing(value) for value in values):
                continue
            subject = "" if is_semantic_missing(values[0]) else str(values[0]).strip()
            measurements = values[1:]
            if all(is_semantic_missing(value) for value in measurements):
                ignored_all_empty_rows += 1
                continue
            if not subject:
                raise ReproductionError(
                    f"Populated legacy tract row has no SubjectID in {path.name}"
                )
            if subject in subjects:
                raise ReproductionError(f"Duplicate SubjectID in {path.name}")
            subjects.add(subject)
            encoded: list[str] = []
            for value in measurements:
                if is_semantic_missing(value):
                    encoded.append("NA")
                    continue
                if isinstance(value, bool) or not isinstance(value, numbers.Real):
                    raise ReproductionError(
                        f"Non-numeric populated tract value in {path.name}"
                    )
                numeric = float(value)
                if not math.isfinite(numeric):
                    raise ReproductionError(f"Non-finite tract value in {path.name}")
                encoded.append(format(numeric, ".5g"))
            lines.append(subject + "\t" + "\t".join(encoded))
        lines.sort(key=lambda line: line.split("\t", 1)[0])
        payload = (
            "semantic_legacy_tract_5sig_v2\n"
            + "\t".join(headers)
            + "\n"
            + "\n".join(lines)
            + "\n"
        ).encode("utf-8")
        return {
            "fingerprint": hashlib.sha256(payload).hexdigest(),
            "populated_subject_rows": len(lines),
            "ignored_all_empty_subject_rows": ignored_all_empty_rows,
            "columns": len(headers),
            "canonical_significant_figures": 5,
        }
    finally:
        workbook.close()


def semantic_noddi_tract_fingerprint(
    path: Path, specification: dict[str, Any]
) -> dict[str, Any]:
    """Fingerprint a NODDI tract CSV after reviewed five-decimal normalization."""
    mode = "semantic_noddi_tract_5dp_v1"
    if str(specification.get("match_mode", mode)) != mode:
        raise ReproductionError(f"Incorrect semantic NODDI mode for {path.name}")
    with path.open("r", encoding="utf-8", newline="") as handle:
        rows = list(csv.reader(handle))
    if not rows:
        raise ReproductionError(f"NODDI tract table is empty: {path.name}")
    header = rows[0]
    if (
        len(header) != 49
        or header[0] != "SubjectID"
        or any(not value for value in header)
        or len(set(header)) != len(header)
    ):
        raise ReproductionError(f"Invalid NODDI tract schema in {path.name}")

    canonical_rows: list[str] = []
    subjects: set[str] = set()
    missing_cells = 0
    for row_number, row in enumerate(rows[1:], start=2):
        if len(row) != len(header):
            raise ReproductionError(
                f"NODDI tract row {row_number} has {len(row)}, not {len(header)}, columns"
            )
        subject = row[0].strip()
        if not subject or subject != row[0] or subject in subjects:
            raise ReproductionError(
                f"Invalid or duplicate SubjectID in {path.name} row {row_number}"
            )
        subjects.add(subject)
        encoded = [subject]
        for raw_value in row[1:]:
            value = raw_value.strip()
            if not value or value.lower() == "nan":
                encoded.append("NA")
                missing_cells += 1
                continue
            numeric = finite_float(value, "tract metric", path.name)
            if not 0 <= numeric <= 1:
                raise ReproductionError(
                    f"NODDI tract value is outside [0,1] in {path.name} row {row_number}"
                )
            encoded.append(format(numeric, ".5f"))
        canonical_rows.append("\t".join(encoded))

    study = analysis_settings()["study"]
    expected_rows = int(study["expected_primary_tract_subjects"])
    expected_missing_cells = int(
        study["expected_noddi_missing_metric_cells_per_table"]
    )
    if len(canonical_rows) != expected_rows or missing_cells != expected_missing_cells:
        raise ReproductionError(
            f"Unexpected NODDI tract cohort/missingness in {path.name}: "
            f"rows={len(canonical_rows)}, missing metric cells={missing_cells}"
        )
    canonical_rows.sort(key=lambda line: line.split("\t", 1)[0])
    payload = (
        mode
        + "\n"
        + "\t".join(header)
        + "\n"
        + "\n".join(canonical_rows)
        + "\n"
    ).encode("utf-8")
    return {
        "fingerprint": hashlib.sha256(payload).hexdigest(),
        "subject_rows": len(canonical_rows),
        "columns": len(header),
        "missing_metric_cells": missing_cells,
        "canonical_decimal_places": 5,
    }


def validate_aggregate_inputs(
    stats_dir: Path, metrics_dir: Path | None = None
) -> dict[str, Any]:
    """Fail closed on the frozen statistics inputs and optional NODDI tables."""
    records: list[dict[str, Any]] = []
    failures: list[str] = []
    settings = analysis_settings()
    for specification in settings["frozen_aggregate_inputs"]:
        directory = str(specification["directory"])
        if directory == "metrics" and metrics_dir is None:
            continue
        name = str(specification["name"])
        generated_metric = (
            metrics_dir / name
            if metrics_dir is not None and name.startswith("GroupTract")
            else None
        )
        if directory == "stats" and generated_metric is not None and generated_metric.is_file():
            # This mirrors the unified MATLAB model: actively regenerated
            # GroupTract workbooks in metricsDir take precedence over the
            # historical reference copies in statsDir.
            root = metrics_dir
        else:
            root = stats_dir if directory == "stats" else metrics_dir
        assert root is not None
        path = resolve_existing(root / name, "aggregate input")
        raw_digest = sha256_file(path)
        mode = str(specification["match_mode"])
        semantic: dict[str, Any] = {}
        if mode == "raw_sha256":
            observed = raw_digest
        elif mode == "semantic_legacy_tract_5sig_v2":
            semantic = semantic_legacy_tract_fingerprint(path)
            observed = str(semantic["fingerprint"])
        elif mode == "semantic_noddi_tract_5dp_v1":
            semantic = semantic_noddi_tract_fingerprint(path, specification)
            observed = str(semantic["fingerprint"])
        else:
            raise ReproductionError(f"Unsupported aggregate-input match mode: {mode}")
        expected_key = (
            "source_expected_fingerprint"
            if metrics_dir is None and "source_expected_fingerprint" in specification
            else "expected_fingerprint"
        )
        expected = str(specification[expected_key])
        match = observed == expected
        if not match:
            failures.append(path.name)
        records.append(
            {
                "name": path.name,
                "source_directory": str(root),
                "match_mode": mode,
                "raw_sha256": raw_digest,
                "observed_fingerprint": observed,
                "expected_fingerprint": expected,
                "matches_completed_run": match,
                "size_bytes": path.stat().st_size,
                **{key: value for key, value in semantic.items() if key != "fingerprint"},
            }
        )
    if failures:
        raise ReproductionError(
            "Frozen aggregate input fingerprint mismatch: " + ", ".join(failures)
        )
    return {"files_checked": len(records), "status": "PASS", "files": records}


MODEL_COLUMNS = [
    "Outcome",
    "SourceField",
    "Metric",
    "Model",
    "Adjusted",
    "Demographics",
    "CasePolicy",
    "N",
    "PositiveN",
    "PerformanceName",
    "Performance",
    "CI_lower",
    "CI_upper",
    "Lambda",
    "p_raw",
    "p_adj",
    "MultiplicityFamilyN",
    "InferenceRole",
    "AIC",
]
OUTCOMES = ("EDSS", "MSPro", "T25FW", "9HPT-D", "9HPT-ND", "SDMT", "MSFC-SDMT")
MRI_METRICS = ("T1", "MTR", "FA", "MD", "NDI", "ODI", "FWF")
NODDI_METRICS = ("NDI", "ODI", "FWF")
MODEL_FRAMEWORKS = ("Tract-based", "Classical-region")


def read_csv_rows(path: Path, expected_columns: Sequence[str], label: str) -> list[dict[str, str]]:
    resolve_existing(path, label)
    with path.open("r", encoding="utf-8", newline="") as handle:
        reader = csv.DictReader(handle)
        if reader.fieldnames != list(expected_columns):
            raise ReproductionError(
                f"Unexpected columns in {label}: {reader.fieldnames}"
            )
        return list(reader)


def parse_adjusted(value: str, label: str) -> int:
    normalized = str(value).strip().lower()
    if normalized in {"0", "false"}:
        return 0
    if normalized in {"1", "true"}:
        return 1
    raise ReproductionError(f"Unexpected Adjusted value in {label}: {value!r}")


def finite_float(value: str, field: str, label: str) -> float:
    try:
        result = float(value)
    except (TypeError, ValueError) as failure:
        raise ReproductionError(f"Non-numeric {field} in {label}: {value!r}") from failure
    if not math.isfinite(result):
        raise ReproductionError(f"Non-finite {field} in {label}: {value!r}")
    return result


def validate_matlab_metric_table_reports(metrics_dir: Path) -> dict[str, Any]:
    """Require either fresh MATLAB reports or the frozen final-table audit."""
    tract_path = metrics_dir / "metric_table_reference_validation.csv"
    classical_path = (
        metrics_dir
        / "ClassicalRegionNODDI_All_Legacy_T1_MTR_FA_MD_validation.csv"
    )
    frozen_audit_path = metrics_dir / "revision_round2_validation.json"
    if frozen_audit_path.is_file() and (
        not tract_path.is_file() or not classical_path.is_file()
    ):
        audit = json.loads(frozen_audit_path.read_text(encoding="utf-8"))
        candidate = audit.get("candidate_table_validation", {})
        if audit.get("status") != "PASS" or candidate.get("Status") != "PASS":
            raise ReproductionError("Frozen final-table audit did not pass")
        expected_missing = {
            "T1": 0,
            "MTR": 96,
            "FA": 0,
            "MD": 0,
            "NDI": 591,
            "ODI": 591,
            "FWF": 591,
        }
        observed: dict[str, dict[str, Any]] = {}
        for row in candidate.get("Rows", []):
            metric = str(row.get("Metric", ""))
            if metric in observed:
                raise ReproductionError(
                    f"Duplicate metric in frozen final-table audit: {metric}"
                )
            observed[metric] = row
        if set(observed) != set(expected_missing):
            raise ReproductionError("Frozen final-table audit has an unexpected metric set")
        for metric, missing_cells in expected_missing.items():
            row = observed[metric]
            if (
                int(row.get("Rows", 0)) != 132
                or int(row.get("MissingCells", -1)) != missing_cells
                or row.get("OriginalRowsUnchanged") is not True
                or row.get("RecoveredRowsMatchMATLABGeneration") is not True
            ):
                raise ReproductionError(
                    f"Frozen final-table audit failed for {metric}"
                )
        classical = audit.get("classical_tables", {})
        for name in ("ClassicalRegionNODDI_All.csv", "ClassicalRegionAllMetrics.csv"):
            record = classical.get(name, {})
            path = resolve_existing(metrics_dir / name, "frozen classical table")
            if (
                int(record.get("rows", 0)) != 87
                or str(record.get("sha256", "")) != sha256_file(path)
            ):
                raise ReproductionError(
                    f"Frozen final-table classical audit failed for {name}"
                )
        return {
            "status": "PASS",
            "mode": "frozen_final_table_audit",
            "tract_metrics_checked": sorted(observed),
            "tract_subjects": 132,
            "classical_subjects": 87,
            "audit_sha256": sha256_file(frozen_audit_path),
        }

    tract_columns = [
        "Metric",
        "ReferenceFile",
        "GeneratedSubjects",
        "ReferenceSubjects",
        "ComparedCells",
        "FailedCells",
        "MissingnessFailures",
        "MaximumAbsoluteError",
        "Passed",
    ]
    tract_rows = read_csv_rows(
        tract_path,
        tract_columns,
        "MATLAB tract reference-validation report",
    )
    rows_by_metric: dict[str, dict[str, str]] = {}
    for row in tract_rows:
        metric = row["Metric"]
        if metric in rows_by_metric:
            raise ReproductionError(
                f"Duplicate metric in MATLAB tract validation report: {metric}"
            )
        rows_by_metric[metric] = row
        if parse_adjusted(row["Passed"], "MATLAB tract validation report") != 1:
            raise ReproductionError(f"MATLAB tract validation failed for {metric}")
        try:
            generated_subjects = int(float(row["GeneratedSubjects"]))
            reference_subjects = int(float(row["ReferenceSubjects"]))
            compared_cells = int(float(row["ComparedCells"]))
            failed_cells = int(float(row["FailedCells"]))
            missingness_failures = int(float(row["MissingnessFailures"]))
        except ValueError as failure:
            raise ReproductionError(
                f"Invalid count in MATLAB tract validation report for {metric}"
            ) from failure
        if (
            generated_subjects < 87
            or reference_subjects < generated_subjects
            or compared_cells <= 0
            or failed_cells != 0
            or missingness_failures != 0
        ):
            raise ReproductionError(
                f"Invalid MATLAB tract validation totals for {metric}"
            )
        finite_float(
            row["MaximumAbsoluteError"],
            "MaximumAbsoluteError",
            f"MATLAB tract validation for {metric}",
        )

    legacy_metrics = {"T1", "MTR", "FA", "MD"}
    all_metrics = legacy_metrics | {"NDI", "ODI", "ISOVF"}
    observed_metrics = set(rows_by_metric)
    if observed_metrics not in (legacy_metrics, all_metrics):
        raise ReproductionError(
            "MATLAB tract validation must contain T1/MTR/FA/MD, with either "
            "all three NODDI reference checks or none"
        )

    classical_columns = [
        "SubjectID",
        "Metric",
        "Region",
        "Calculated",
        "Expected",
        "AbsoluteError",
        "Passed",
        "Tolerance",
    ]
    classical_rows = read_csv_rows(
        classical_path,
        classical_columns,
        "MATLAB classical-region cellwise validation report",
    )
    expected_regions = {
        "periventricular",
        "juxtacortical",
        "infratentorial",
        "deepwhitematter",
    }
    expected_row_count = 87 * len(legacy_metrics) * len(expected_regions)
    if len(classical_rows) != expected_row_count:
        raise ReproductionError(
            "Expected 1,392 MATLAB classical-region validation rows, found "
            f"{len(classical_rows)}"
        )
    observed_keys: set[tuple[str, str, str]] = set()
    subjects: set[str] = set()
    maximum_error = 0.0
    for row in classical_rows:
        key = (row["SubjectID"], row["Metric"], row["Region"])
        if key in observed_keys:
            raise ReproductionError(
                "Duplicate subject/metric/region in MATLAB classical validation report"
            )
        observed_keys.add(key)
        subjects.add(row["SubjectID"])
        if row["Metric"] not in legacy_metrics or row["Region"] not in expected_regions:
            raise ReproductionError(
                "Unexpected metric or region in MATLAB classical validation report"
            )
        if parse_adjusted(row["Passed"], "MATLAB classical validation report") != 1:
            raise ReproductionError("MATLAB classical-region cellwise validation failed")
        absolute_error = finite_float(
            row["AbsoluteError"],
            "AbsoluteError",
            "MATLAB classical validation report",
        )
        tolerance = finite_float(
            row["Tolerance"],
            "Tolerance",
            "MATLAB classical validation report",
        )
        if absolute_error > tolerance or tolerance > 1e-12:
            raise ReproductionError(
                "MATLAB classical validation exceeds its locked 1e-12 tolerance"
            )
        maximum_error = max(maximum_error, absolute_error)
    if len(subjects) != 87:
        raise ReproductionError(
            f"Expected 87 subjects in MATLAB classical validation, found {len(subjects)}"
        )

    return {
        "status": "PASS",
        "tract_metrics_checked": sorted(observed_metrics),
        "tract_reference_rows": len(tract_rows),
        "classical_subjects": len(subjects),
        "classical_cellwise_rows": len(classical_rows),
        "classical_maximum_absolute_error": maximum_error,
        "tract_report_sha256": sha256_file(tract_path),
        "classical_report_sha256": sha256_file(classical_path),
    }


def validate_model_outputs(
    all_models_path: Path,
    noddi_compatibility_path: Path,
    full_figure_path: Path,
    matched_figure_path: Path,
) -> dict[str, Any]:
    """Validate the active 196-row analysis and both 70-row figure inputs."""
    all_rows = read_csv_rows(all_models_path, MODEL_COLUMNS, "all-qMRI model table")
    if len(all_rows) != 196:
        raise ReproductionError(
            f"Expected 196 active all-qMRI model rows, found {len(all_rows)}"
        )

    expected_keys = {
        (outcome, metric, framework, adjusted)
        for outcome in OUTCOMES
        for metric in MRI_METRICS
        for framework in MODEL_FRAMEWORKS
        for adjusted in (0, 1)
    }
    rows_by_key: dict[tuple[str, str, str, int], dict[str, str]] = {}
    significant_adjusted = 0
    for row in all_rows:
        adjusted = parse_adjusted(row["Adjusted"], "all-qMRI model table")
        key = (row["Outcome"], row["Metric"], row["Model"], adjusted)
        if key in rows_by_key:
            raise ReproductionError(f"Duplicate all-qMRI model row: {key}")
        rows_by_key[key] = row
        try:
            sample_size = int(row["N"])
            family_size = int(row["MultiplicityFamilyN"])
        except ValueError as failure:
            raise ReproductionError(f"Invalid integer field in all-qMRI row {key}") from failure
        if sample_size < 20:
            raise ReproductionError(f"Implausible model sample size in row {key}: {sample_size}")
        p_raw = finite_float(row["p_raw"], "p_raw", f"all-qMRI row {key}")
        if not 0 <= p_raw <= 1:
            raise ReproductionError(f"p_raw is outside [0,1] in row {key}")
        finite_float(row["Performance"], "Performance", f"all-qMRI row {key}")
        finite_float(row["Lambda"], "Lambda", f"all-qMRI row {key}")
        finite_float(row["AIC"], "AIC", f"all-qMRI row {key}")
        if adjusted:
            p_adjusted = finite_float(row["p_adj"], "p_adj", f"all-qMRI row {key}")
            expected_adjusted = min(p_raw * 98, 1.0)
            if family_size != 98 or not math.isclose(
                p_adjusted, expected_adjusted, rel_tol=0, abs_tol=1e-12
            ):
                raise ReproductionError(
                    f"Incorrect m=98 multiplicity result in all-qMRI row {key}"
                )
            significant_adjusted += int(p_adjusted < 0.05)
        else:
            p_adjusted_text = str(row["p_adj"]).strip().lower()
            if family_size != 0 or p_adjusted_text not in {"", "nan"}:
                raise ReproductionError(
                    f"No-demographic row unexpectedly entered a p family: {key}"
                )
    if set(rows_by_key) != expected_keys:
        missing = sorted(expected_keys - set(rows_by_key))
        extra = sorted(set(rows_by_key) - expected_keys)
        raise ReproductionError(
            f"All-qMRI Cartesian grid mismatch; missing={missing[:5]}, extra={extra[:5]}"
        )
    if significant_adjusted != 46:
        raise ReproductionError(
            f"Expected 46 significant adjusted rows at m=98, found {significant_adjusted}"
        )

    noddi_rows = read_csv_rows(
        noddi_compatibility_path,
        MODEL_COLUMNS,
        "84-row NODDI compatibility model table",
    )
    if len(noddi_rows) != 84:
        raise ReproductionError(
            f"Expected 84 NODDI compatibility rows, found {len(noddi_rows)}"
        )
    noddi_keys: set[tuple[str, str, str, int]] = set()
    for row in noddi_rows:
        key = (
            row["Outcome"],
            row["Metric"],
            row["Model"],
            parse_adjusted(row["Adjusted"], "NODDI compatibility model table"),
        )
        if key in noddi_keys or key not in rows_by_key or key[1] not in NODDI_METRICS:
            raise ReproductionError(f"Invalid NODDI compatibility model row: {key}")
        noddi_keys.add(key)
        if row != rows_by_key[key]:
            raise ReproductionError(
                f"NODDI compatibility row differs from the unified result: {key}"
            )
    expected_noddi_keys = {key for key in expected_keys if key[1] in NODDI_METRICS}
    if noddi_keys != expected_noddi_keys:
        raise ReproductionError("NODDI compatibility table is not the exact unified subset")

    figure_reports: dict[str, dict[str, Any]] = {}
    figure_columns = ["ClinicalMetric", "Model", "N", "R2", "R2_Demo"]
    figure_outcomes = ("T25FW", "x9HPTD", "x9HPTND", "SDMTcorrect", "MSFC_SDMT")
    figure_metrics = ("T1", "MTR", "FA", "MD", "NDI", "ODI", "ISOVF")
    figure_frameworks = ("tract-based", "classic")
    expected_figure_keys = {
        (f"{outcome}-{metric}", framework)
        for outcome in figure_outcomes
        for metric in figure_metrics
        for framework in figure_frameworks
    }
    for label, path in (
        ("full-sample figure table", full_figure_path),
        ("matched-subject figure table", matched_figure_path),
    ):
        rows = read_csv_rows(path, figure_columns, label)
        if len(rows) != 70:
            raise ReproductionError(f"Expected 70 rows in {label}, found {len(rows)}")
        observed: set[tuple[str, str]] = set()
        for row in rows:
            key = (row["ClinicalMetric"], row["Model"])
            if key in observed:
                raise ReproductionError(f"Duplicate row in {label}: {key}")
            observed.add(key)
            try:
                sample_size = int(row["N"])
            except ValueError as failure:
                raise ReproductionError(f"Invalid N in {label}: {key}") from failure
            if sample_size < 20:
                raise ReproductionError(f"Implausible N in {label}: {key}")
            finite_float(row["R2"], "R2", label)
            finite_float(row["R2_Demo"], "R2_Demo", label)
        if observed != expected_figure_keys:
            raise ReproductionError(f"Cartesian grid mismatch in {label}")
        figure_reports[path.name] = {
            "rows": len(rows),
            "sha256": sha256_file(path),
        }

    return {
        "status": "PASS",
        "all_qmri_rows": len(all_rows),
        "adjusted_rows": 98,
        "no_demographic_rows": 98,
        "significant_adjusted_rows_m98": significant_adjusted,
        "noddi_compatibility_rows": len(noddi_rows),
        "all_qmri_sha256": sha256_file(all_models_path),
        "noddi_compatibility_sha256": sha256_file(noddi_compatibility_path),
        "figure_tables": figure_reports,
    }


def validate_manuscript_table_outputs(table_dir: Path) -> dict[str, Any]:
    """Validate the five plain manuscript CSV grids."""
    expected_shapes = {
        "Main_Table_1.csv": (13, 5),
        "Main_Table_2.csv": (15, 5),
        "Supplementary_Table_3.csv": (15, 5),
        "Supplementary_Table_4.csv": (99, 6),
        "Supplementary_Table_5.csv": (58, 6),
    }
    records: dict[str, dict[str, Any]] = {}
    loaded: dict[str, list[list[str]]] = {}
    for filename, expected_shape in expected_shapes.items():
        path = resolve_existing(table_dir / filename, "manuscript CSV table")
        with path.open("r", encoding="utf-8", newline="") as handle:
            rows = list(csv.reader(handle))
        observed_shape = (len(rows), max((len(row) for row in rows), default=0))
        if observed_shape != expected_shape or any(
            len(row) != expected_shape[1] for row in rows
        ):
            raise ReproductionError(
                f"Unexpected grid for {filename}: {observed_shape}; expected {expected_shape}"
            )
        loaded[filename] = rows
        records[filename] = {
            "rows": observed_shape[0],
            "columns": observed_shape[1],
            "sha256": sha256_file(path),
        }

    if loaded["Main_Table_1.csv"][0] != [
        "Parameter", "Healthy Controls", "RRMS", "PPMS", "SPMS"
    ]:
        raise ReproductionError("Main Table 1 header changed")
    if loaded["Supplementary_Table_4.csv"][0] != [
        "Outcome", "MRI Metric", "Model", "N", "p", "p_adj"
    ]:
        raise ReproductionError("Supplementary Table 4 header changed")
    if loaded["Supplementary_Table_5.csv"][17] != [
        "Outcome", "Metric", "Model", "N", "R2", "AIC"
    ]:
        raise ReproductionError("Supplementary Table 5 continuous header changed")
    return {"status": "PASS", "files": records}


def copy_verified_local_input(source: Path, destination: Path, force: bool) -> dict[str, Any]:
    """Copy a frozen input into the local active-table directory without ambiguity."""
    resolve_existing(source, "snapshot analysis input")
    destination.parent.mkdir(parents=True, exist_ok=True)
    source_hash = sha256_file(source)
    if destination.exists():
        if not destination.is_file():
            raise ReproductionError(f"Analysis-input destination is not a file: {destination}")
        destination_hash = sha256_file(destination)
        if destination_hash == source_hash:
            return {"name": destination.name, "sha256": source_hash, "status": "REUSED"}
        if not force:
            raise ReproductionError(
                f"Local analysis input already exists with different content: {destination}. "
                "Pass --force to replace generated local outputs."
            )
    temporary = destination.with_suffix(destination.suffix + ".partial")
    if temporary.exists():
        temporary.unlink()
    shutil.copy2(source, temporary)
    if sha256_file(temporary) != source_hash:
        temporary.unlink(missing_ok=True)
        raise ReproductionError(f"Copied analysis input failed verification: {destination.name}")
    os.replace(temporary, destination)
    return {"name": destination.name, "sha256": source_hash, "status": "COPIED"}


def verify_multiplicity_accounting() -> dict[str, Any]:
    """Recompute the 56-to-98 family accounting from the bundled 196-row result."""
    bundled_path = (
        PACKAGE_ROOT / "results" / "model_results" / "all_qmri_manuscript_models.csv"
    )
    all_rows = read_csv_rows(
        bundled_path,
        MODEL_COLUMNS,
        "bundled all-qMRI model result",
    )
    if len(all_rows) != 196:
        raise ReproductionError(
            f"Expected 196 rows in the bundled all-qMRI result, found {len(all_rows)}"
        )

    adjusted_rows: list[dict[str, str]] = []
    adjusted_keys: set[tuple[str, str, str]] = set()
    for row in all_rows:
        if parse_adjusted(row["Adjusted"], "bundled all-qMRI model result") != 1:
            continue
        key = (row["Outcome"], row["Model"], row["Metric"])
        if key in adjusted_keys:
            raise ReproductionError(f"Duplicate adjusted row in bundled result: {key}")
        adjusted_keys.add(key)
        raw = finite_float(row["p_raw"], "p_raw", "bundled all-qMRI model result")
        adjusted_98 = finite_float(
            row["p_adj"], "p_adj", "bundled all-qMRI model result"
        )
        try:
            family_size = int(row["MultiplicityFamilyN"])
        except ValueError as failure:
            raise ReproductionError(
                f"Invalid multiplicity-family size in bundled result row {key}"
            ) from failure
        if (
            not 0 <= raw <= 1
            or family_size != 98
            or not math.isclose(
                adjusted_98, min(raw * 98, 1.0), rel_tol=0, abs_tol=1e-12
            )
        ):
            raise ReproductionError(
                f"Invalid m=98 multiplicity value in bundled result row {key}"
            )
        adjusted_rows.append(row)

    legacy_metrics = {"T1", "MTR", "FA", "MD"}
    noddi_metrics = set(NODDI_METRICS)
    legacy = [row for row in adjusted_rows if row["Metric"] in legacy_metrics]
    noddi = [row for row in adjusted_rows if row["Metric"] in noddi_metrics]
    unexpected_metrics = {
        row["Metric"] for row in adjusted_rows
    } - legacy_metrics - noddi_metrics
    if len(adjusted_rows) != 98 or len(legacy) != 56 or len(noddi) != 42:
        raise ReproductionError(
            "Bundled adjusted model family must contain 56 legacy and 42 NODDI rows; "
            f"found total={len(adjusted_rows)}, legacy={len(legacy)}, NODDI={len(noddi)}"
        )
    if unexpected_metrics:
        raise ReproductionError(
            f"Unexpected adjusted metrics in bundled result: {sorted(unexpected_metrics)}"
        )

    significant_56 = 0
    significant_98 = 0
    lost: list[dict[str, str]] = []
    for row in legacy:
        raw = finite_float(row["p_raw"], "p_raw", "bundled legacy model result")
        adjusted_56 = min(raw * 56, 1.0)
        adjusted_98 = min(raw * 98, 1.0)
        passes_56 = adjusted_56 < 0.05
        passes_98 = adjusted_98 < 0.05
        significant_56 += int(passes_56)
        significant_98 += int(passes_98)
        if passes_56 and not passes_98:
            lost.append(
                {"model": row["Model"], "metric": row["Metric"], "outcome": row["Outcome"]}
            )

    frameworks = {row["Model"] for row in noddi}
    if frameworks != {"Tract-based", "Classical-region"}:
        raise ReproductionError(f"Unexpected NODDI model frameworks: {sorted(frameworks)}")
    noddi_significant = sum(
        min(finite_float(row["p_raw"], "p_raw", "bundled NODDI model result") * 98, 1.0)
        < 0.05
        for row in noddi
    )

    expected_lost = [
        {"model": "Classical-region", "metric": "MTR", "outcome": "EDSS"},
        {"model": "Classical-region", "metric": "MD", "outcome": "EDSS"},
        {"model": "Tract-based", "metric": "MD", "outcome": "MSPro"},
    ]
    lost_keys = {(row["outcome"], row["model"], row["metric"]) for row in lost}
    expected_lost_keys = {
        (row["outcome"], row["model"], row["metric"]) for row in expected_lost
    }
    if lost_keys != expected_lost_keys:
        raise ReproductionError(
            "Multiplicity summary mismatch for previous rows losing significance at m=98"
        )

    expected = {
        "previous_family_size": 56,
        "previous_significant_at_m56": 30,
        "previous_significant_after_expansion_to_m98": 27,
        "noddi_tests_added": 42,
        "noddi_significant_at_m98": 19,
        "combined_family_size": 98,
        "combined_significant_at_m98": 46,
    }
    observed = {
        "previous_family_size": len(legacy),
        "previous_significant_at_m56": significant_56,
        "previous_significant_after_expansion_to_m98": significant_98,
        "previous_rows_losing_significance_at_m98": lost,
        "noddi_tests_added": len(noddi),
        "noddi_significant_at_m98": noddi_significant,
        "combined_family_size": len(legacy) + len(noddi),
        "combined_significant_at_m98": significant_98 + noddi_significant,
    }
    for key, value in observed.items():
        if key == "previous_rows_losing_significance_at_m98":
            continue
        if expected.get(key) != value:
            raise ReproductionError(f"Multiplicity summary mismatch for {key}")
    return {
        **observed,
        "bundled_result_rows": len(all_rows),
        "bundled_result_sha256": sha256_file(bundled_path),
        "status": "PASS",
    }


def validate_fit_cohort(fit_root: Path, bids_root: Path) -> dict[str, Any]:
    fit_subject_root = fit_root / "derivatives" / "ses-01"
    if not fit_subject_root.is_dir():
        raise ReproductionError(f"NODDI fit subject directory is missing: {fit_subject_root}")
    observed = sorted(path.name for path in fit_subject_root.glob("sub-*") if path.is_dir())
    source_root = bids_root / "derivatives" / "TractoFlow" / "ses-01"
    expected = sorted(path.name for path in source_root.glob("sub-*") if path.is_dir())
    if observed != expected:
        raise ReproductionError(
            "NODDI fit cohort does not exactly match the validated source cohort; "
            "participant labels are intentionally not printed by the public runner."
        )
    status_path = fit_root / "logs" / "run_status.csv"
    if not status_path.is_file():
        raise ReproductionError(f"NODDI fit status file is missing: {status_path}")
    with status_path.open("r", encoding="utf-8", newline="") as handle:
        status_rows = list(csv.DictReader(handle))
    status_by_subject = {str(row.get("subject_id", "")): row for row in status_rows}
    if len(status_by_subject) != len(status_rows) or sorted(status_by_subject) != expected:
        raise ReproductionError("NODDI fit status cohort does not match the source cohort")

    config_path = fit_root / "config" / "full_cohort_config.json"
    if not config_path.is_file():
        raise ReproductionError(f"NODDI fit configuration is missing: {config_path}")
    fit_config = json.loads(config_path.read_text(encoding="utf-8"))
    if Path(fit_config.get("source_root", "")).resolve() != source_root.resolve():
        raise ReproductionError("NODDI fit configuration refers to a different source tree")
    if Path(fit_config.get("output_root", "")).resolve() != fit_root.resolve():
        raise ReproductionError("NODDI fit configuration refers to a different output tree")
    expected_subject_file = Path(fit_config.get("expected_source_subjects_file", ""))
    if not expected_subject_file.is_file():
        raise ReproductionError("NODDI fit configuration lacks its private cohort file")
    configured_subjects = sorted(
        line.strip()
        for line in expected_subject_file.read_text(encoding="utf-8").splitlines()
        if line.strip()
    )
    if configured_subjects != expected:
        raise ReproductionError("NODDI fit private cohort file does not match the source cohort")
    settings = analysis_settings()
    expected_model = dict(settings["model"])
    expected_model.pop("implementation", None)
    if (
        fit_config.get("session") != settings["study"]["session"]
        or fit_config.get("acquisition") != settings["acquisition"]
        or fit_config.get("model") != expected_model
        or fit_config.get("evaluation") != settings["evaluation"]
    ):
        raise ReproductionError("NODDI fit scientific configuration is not the frozen configuration")
    execution = fit_config.get("execution", {})
    if (
        execution.get("bootstrap_subject") != expected[0]
        or execution.get("resume_completed_subjects") is not True
        or int(execution.get("default_subject_processes", 0)) not in range(1, 5)
    ):
        raise ReproductionError("NODDI fit execution configuration failed validation")
    config_digest = sha256_file(config_path)
    required_outputs = (
        "fit_NDI.nii.gz",
        "fit_ODI.nii.gz",
        "fit_FWF.nii.gz",
        "fit_dir.nii.gz",
        "fit_RMSE.nii.gz",
        "fit_NRMSE.nii.gz",
        "run_config.json",
        "subject_metrics.json",
    )
    for subject in expected:
        status = status_by_subject[subject]
        if status.get("fit_status") not in {"SUCCESS", "SKIPPED_COMPLETE"}:
            raise ReproductionError("The completed NODDI fit contains a failed subject")
        if status.get("technical_qc") != "PASS":
            raise ReproductionError("The completed NODDI fit contains a technical-QC failure")
        subject_root = fit_subject_root / subject
        if any(
            not (subject_root / name).is_file() or (subject_root / name).stat().st_size == 0
            for name in required_outputs
        ):
            raise ReproductionError("The completed NODDI fit is missing required subject outputs")
        run_config = json.loads((subject_root / "run_config.json").read_text(encoding="utf-8"))
        metrics = json.loads((subject_root / "subject_metrics.json").read_text(encoding="utf-8"))
        if (
            run_config.get("status") != "SUCCESS"
            or run_config.get("subject_id") != subject
            or run_config.get("analysis_config_sha256") != config_digest
            or metrics.get("fit_status") != "SUCCESS"
            or metrics.get("technical_qc") != "PASS"
        ):
            raise ReproductionError("The completed NODDI fit metadata failed validation")
        source_inputs = {
            "dwi": source_root / subject / "Normalize_DWI" / f"{subject}__dwi_normalized.nii.gz",
            "mask": source_root / subject / "Crop_DWI" / f"{subject}__b0_mask_cropped.nii.gz",
            "bval": source_root / subject / "Eddy_Topup" / f"{subject}__bval_eddy",
            "bvec": source_root / subject / "Eddy_Topup" / f"{subject}__dwi_eddy_corrected.bvec",
        }
        for label, source in source_inputs.items():
            recorded = run_config.get("inputs", {}).get(label, {})
            stat = source.stat()
            current: dict[str, Any] = {
                "path": str(source),
                "size_bytes": int(stat.st_size),
                "mtime_ns": int(stat.st_mtime_ns),
                "sha256": sha256_file(source),
            }
            if recorded != current:
                raise ReproductionError("A completed NODDI fit input no longer matches its source")
        for name in FIT_HASHED_OUTPUTS:
            output = subject_root / name
            current_output = {
                "size_bytes": int(output.stat().st_size),
                "sha256": sha256_file(output),
            }
            if run_config.get("outputs", {}).get(name) != current_output:
                raise ReproductionError("A completed NODDI fit output failed content validation")
    return {
        "subjects": len(observed),
        "matches_validated_source_cohort": True,
        "run_status_all_pass": True,
        "per_subject_fit_metadata_matches_source": True,
        "status": "PASS",
    }


def locate_executable(requested: Path | None, names: Sequence[str], label: str) -> Path:
    if requested is not None:
        candidate = requested.expanduser().absolute()
        if candidate.is_file() and os.access(candidate, os.X_OK):
            return candidate
        raise ReproductionError(f"{label} is not executable: {candidate}")
    for name in names:
        found = shutil.which(name)
        if found:
            return Path(found).absolute()
        candidate = Path(name).expanduser()
        if candidate.is_file() and os.access(candidate, os.X_OK):
            return candidate.absolute()
    raise ReproductionError(f"Could not find {label}; pass its path explicitly")


def locate_freesurfer_matlab(requested: Path | None) -> Path:
    """Locate the FreeSurfer MATLAB utilities and require MRIread.m."""
    if requested is not None:
        candidate = requested.expanduser().absolute()
        if candidate.is_dir() and (candidate / "MRIread.m").is_file():
            return candidate
        raise ReproductionError(
            "FreeSurfer MATLAB utilities directory must exist and contain MRIread.m: "
            f"{candidate}"
        )

    candidates: list[Path] = []
    freesurfer_home = os.environ.get("FREESURFER_HOME")
    if freesurfer_home:
        candidates.append(Path(freesurfer_home).expanduser() / "matlab")
    candidates.append(Path("/Applications/freesurfer/7.4.1/matlab"))
    for candidate in candidates:
        candidate = candidate.absolute()
        if candidate.is_dir() and (candidate / "MRIread.m").is_file():
            return candidate
    raise ReproductionError(
        "FreeSurfer MATLAB utilities with MRIread.m were not found; pass "
        "--freesurfer-matlab /path/to/freesurfer/matlab or set FREESURFER_HOME"
    )


def dependency_report(python: Path) -> dict[str, Any]:
    required = {
        "dmri-amico": "2.1.1",
        "dmri-dicelib": "1.1.2",
        "dipy": "1.9.0",
        "nibabel": "5.2.1",
        "numpy": "1.26.4",
        "openpyxl": "3.1.5",
        "matplotlib": "3.8.4",
        "scipy": "1.12.0",
        "threadpoolctl": "3.5.0",
    }
    probe = (
        "import importlib.metadata, json, sys\n"
        f"names = {list(required)!r}\n"
        "versions = {}\n"
        "for n in names:\n"
        " try: versions[n]=importlib.metadata.version(n)\n"
        " except importlib.metadata.PackageNotFoundError: versions[n]='NOT INSTALLED'\n"
        "print(json.dumps({'python':sys.version.split()[0],"
        "'executable':sys.executable,'packages':versions},sort_keys=True))"
    )
    completed = subprocess.run(
        [str(python), "-c", probe],
        text=True,
        capture_output=True,
        check=False,
    )
    if completed.returncode != 0:
        raise ReproductionError(
            f"Could not inspect Python environment {python}: {completed.stderr.strip()}"
        )
    inspected = json.loads(completed.stdout)
    installed = inspected["packages"]
    mismatches: dict[str, dict[str, str]] = {}
    if not str(inspected["python"]).startswith("3.12."):
        mismatches["python"] = {
            "expected": "3.12.x",
            "actual": str(inspected["python"]),
        }
    for package, expected in required.items():
        actual = installed[package]
        if actual != expected:
            mismatches[package] = {"expected": expected, "actual": actual}
    return {
        "python": inspected["python"],
        "executable": inspected["executable"],
        "packages": installed,
        "mismatches": mismatches,
    }


def build_fit_config(bids_root: Path, fit_root: Path, workers: int) -> Path:
    settings = analysis_settings()
    source_subject_root = bids_root / "derivatives" / "TractoFlow" / "ses-01"
    subjects = sorted(path.name for path in source_subject_root.glob("sub-*") if path.is_dir())
    private_subject_file = fit_root / "config" / "expected_source_subjects.txt"
    atomic_text(private_subject_file, canonical_subject_payload(subjects).decode("utf-8"))
    execution = dict(settings["execution"])
    selection = execution.pop("bootstrap_subject_selection", "")
    if selection != "lexicographically_first_source_subject":
        raise ReproductionError(f"Unsupported bootstrap subject rule: {selection!r}")
    execution["bootstrap_subject"] = subjects[0]
    execution["default_subject_processes"] = workers
    configuration = {
        "run_name": "MS100 AMICO-NODDI full ses-01 source-cohort run",
        "session": settings["study"]["session"],
        "source_root": str(bids_root / "derivatives" / "TractoFlow" / "ses-01"),
        "output_root": str(fit_root),
        "expected_source_subjects_file": str(private_subject_file),
        "group_code_mapping": {"0": "Control", "1": "RRMS", "2": "PPMS", "3": "SPMS"},
        "acquisition": settings["acquisition"],
        "model": settings["model"] | {"name": settings["model"]["name"]},
        "evaluation": settings["evaluation"],
        "execution": execution,
    }
    configuration["model"].pop("implementation", None)
    path = fit_root / "config" / "full_cohort_config.json"
    atomic_json(path, configuration)
    return path


def verify_snapshot_manifest(tract_root: Path) -> dict[str, Any]:
    manifest_path = tract_root / "logs" / "source_snapshot_manifest.csv"
    if not manifest_path.is_file():
        raise ReproductionError(f"Snapshot manifest is missing: {manifest_path}")
    rows = list(csv.DictReader(manifest_path.open("r", encoding="utf-8", newline="")))
    failures: list[str] = []
    bytes_verified = 0
    for row in rows:
        if row.get("status") not in {"COPIED", "UPDATED", "REUSED"}:
            if row.get("required", "").lower() in {"true", "1"}:
                failures.append(f"{row.get('status')}: {row.get('source')}")
            continue
        destination = Path(row["destination"])
        expected = row.get("destination_sha256", "")
        if not destination.is_file() or not expected or sha256_file(destination) != expected:
            failures.append(f"hash mismatch: {destination}")
        else:
            bytes_verified += destination.stat().st_size
    if failures:
        raise ReproductionError("Snapshot verification failed:\n  " + "\n  ".join(failures[:30]))
    return {"files": len(rows), "bytes_verified": bytes_verified, "status": "PASS"}


def matlab_literal(path: Path) -> str:
    return str(path).replace("'", "''")


def run_fit(args: argparse.Namespace, paths: dict[str, Path], manifest: dict[str, Any]) -> None:
    config = (
        paths["fit"] / "config" / "full_cohort_config.json"
        if args.dry_run
        else build_fit_config(paths["bids"], paths["fit"], args.workers)
    )
    environment = os.environ.copy()
    environment["MS100_NODDI_WORK_ROOT"] = str(paths["fit"])
    environment["MS100_NODDI_CONFIG"] = str(config)
    command = [
        str(paths["python"]),
        str(PIPELINE_ROOT / "fit_noddi_full_cohort.py"),
        "--max-workers",
        str(args.workers),
    ]
    run_command(
        command,
        manifest=manifest,
        stage="fit_noddi",
        env=environment,
        dry_run=args.dry_run,
    )


def run_extract(args: argparse.Namespace, paths: dict[str, Path], manifest: dict[str, Any]) -> None:
    if not args.dry_run:
        manifest["fit_cohort_validation"] = validate_fit_cohort(paths["fit"], paths["bids"])
    snapshot_command = [
        str(paths["python"]),
        str(PIPELINE_ROOT / "snapshot_inputs.py"),
        "--bids-root",
        str(paths["bids"]),
        "--noddi-root",
        str(paths["fit"]),
        "--output-root",
        str(paths["tract"]),
        "--workers",
        str(args.snapshot_workers),
    ]
    if paths.get("group_mask_overlay") is not None:
        snapshot_command.extend(
            ["--group-mask-overlay", str(paths["group_mask_overlay"])]
        )
    run_command(
        snapshot_command,
        manifest=manifest,
        stage="snapshot_inputs",
        dry_run=args.dry_run,
    )
    if not args.dry_run:
        manifest["snapshot_verification"] = verify_snapshot_manifest(paths["tract"])

    environment = os.environ.copy()
    environment["MS100_NODDI_TRACT_WORK_ROOT"] = str(paths["tract"])
    environment["MS100_NODDI_FIT_ROOT"] = str(paths["fit"])
    environment["MS100_ANTS_APPLY_TRANSFORMS"] = str(paths["ants"])
    extract_command = [
        str(paths["python"]),
        str(PIPELINE_ROOT / "extract_noddi_tract_metrics.py"),
        "--snapshot-root",
        str(paths["tract"] / "source_snapshot"),
        "--noddi-root",
        str(paths["fit"]),
        "--output-root",
        str(paths["tract"]),
        "--ants",
        str(paths["ants"]),
        "--mode",
        "maps-only",
    ]
    if args.qc_figures:
        extract_command.append("--make-qc-figures")
    if args.force:
        extract_command.append("--force")
    run_command(
        extract_command,
        manifest=manifest,
        stage="register_and_qc_noddi_maps",
        env=environment,
        dry_run=args.dry_run,
    )
    if not MATLAB_METRIC_TABLE_RUNNER.is_file():
        raise ReproductionError(
            f"MATLAB all-metric table runner is missing: {MATLAB_METRIC_TABLE_RUNNER}"
        )
    if paths.get("freesurfer_matlab") is None:
        raise ReproductionError(
            "FreeSurfer MATLAB utilities are required for all-metric aggregation"
        )
    snapshot_root = paths["tract"] / "source_snapshot"
    snapshot_stats = snapshot_root / "derivatives" / "derivativesStats"
    metric_output = paths["tract"] / "results"
    metric_expression = (
        f"addpath('{matlab_literal(MATLAB_METRIC_TABLE_RUNNER.parent)}','-begin'); "
        f"generate_all_metric_tables("
        f"'SourceRoot','{matlab_literal(snapshot_root)}',"
        f"'NODDIMapsDir','{matlab_literal(paths['tract'] / 'transformed_maps')}',"
        f"'OutputDir','{matlab_literal(metric_output)}',"
        f"'LegacyStatsDir','{matlab_literal(snapshot_stats)}',"
        f"'FreesurferMatlabDir','{matlab_literal(paths['freesurfer_matlab'])}',"
        f"'FreesurferDir','{matlab_literal(snapshot_root / 'derivatives' / 'freesurfer')}',"
        f"'Force',{'true' if args.force else 'false'});"
    )
    run_command(
        [str(paths["matlab"]), "-batch", metric_expression],
        manifest=manifest,
        stage="aggregate_all_metric_tables_matlab",
        dry_run=args.dry_run,
    )
    if not args.dry_run:
        # The unified MATLAB model accepts one stats directory. Assemble it
        # locally from the actively regenerated tract tables plus immutable
        # copies of the clinical and legacy classical source workbooks.
        copied_inputs = []
        for filename in ("clinicalScore.xlsx", "icometrixLesionStats.xlsx"):
            copied_inputs.append(
                copy_verified_local_input(
                    snapshot_stats / filename,
                    metric_output / filename,
                    args.force,
                )
            )
        expected_tables = (
            "GroupTractT1_All.xlsx",
            "GroupTractT1_All.csv",
            "GroupTractMTR_All.xlsx",
            "GroupTractMTR_All.csv",
            "GroupTractFA_All.xlsx",
            "GroupTractFA_All.csv",
            "GroupTractMD_All.xlsx",
            "GroupTractMD_All.csv",
            "GroupTractNDI_All.csv",
            "GroupTractODI_All.csv",
            "GroupTractFWF_All.csv",
            "ClassicalRegionNODDI_All.csv",
            "ClassicalRegionAllMetrics.csv",
            "metric_table_generation_QC.csv",
            "metric_table_reference_validation.csv",
            "ClassicalRegionNODDI_All_QC.csv",
            "ClassicalRegionNODDI_All_Legacy_T1_MTR_FA_MD_validation.csv",
            "ClassicalRegionNODDI_All_Legacy_FA_MD_validation.csv",
        )
        table_records: list[dict[str, Any]] = []
        for filename in expected_tables:
            path = resolve_existing(metric_output / filename, f"MATLAB metric output {filename}")
            table_records.append(
                {"name": filename, "size_bytes": path.stat().st_size, "sha256": sha256_file(path)}
            )
        manifest["matlab_metric_table_generation"] = {
            "tables": table_records,
            "copied_frozen_inputs": copied_inputs,
            "status": "PASS",
        }
        manifest["matlab_metric_table_validation"] = (
            validate_matlab_metric_table_reports(metric_output)
        )
        manifest["extracted_aggregate_validation"] = validate_aggregate_inputs(
            metric_output,
            metric_output,
        )


def run_stats(args: argparse.Namespace, paths: dict[str, Path], manifest: dict[str, Any]) -> None:
    if not MATLAB_RUNNER.is_file():
        raise ReproductionError(f"MATLAB statistics runner is missing: {MATLAB_RUNNER}")
    if not MATLAB_FIGURE_RUNNER.is_file():
        raise ReproductionError(f"MATLAB figure runner is missing: {MATLAB_FIGURE_RUNNER}")
    if not MATLAB_TABLE_RUNNER.is_file():
        raise ReproductionError(f"MATLAB table runner is missing: {MATLAB_TABLE_RUNNER}")
    stats_dir = (
        args.stats_dir.expanduser().resolve()
        if args.stats_dir
        else paths["tract"] / "results"
    )
    metrics_dir = (
        args.metrics_dir.expanduser().resolve()
        if args.metrics_dir
        else paths["tract"] / "results"
    )
    classical_file = (
        args.classical_file.expanduser().resolve()
        if args.classical_file
        else metrics_dir / "ClassicalRegionNODDI_All.csv"
    )
    required_stats = (
        "clinicalScore.xlsx",
        "icometrixLesionStats.xlsx",
    )
    required_metrics = (
        "GroupTractT1_All.xlsx",
        "GroupTractMTR_All.xlsx",
        "GroupTractFA_All.xlsx",
        "GroupTractMD_All.xlsx",
        "GroupTractNDI_All.csv",
        "GroupTractODI_All.csv",
        "GroupTractFWF_All.csv",
    )
    if not args.dry_run or args.stage == "stats":
        for filename in required_stats:
            resolve_existing(stats_dir / filename, f"statistics input {filename}")
        for filename in required_metrics:
            resolve_existing(metrics_dir / filename, f"qMRI metric table {filename}")
        resolve_existing(classical_file, "classical-region NODDI table")
        combined_classical_file = resolve_existing(
            metrics_dir / "ClassicalRegionAllMetrics.csv",
            "combined seven-metric classical-region table",
        )
        manifest["statistics_input_validation"] = validate_aggregate_inputs(
            stats_dir, metrics_dir
        )
        manifest["matlab_metric_table_validation"] = (
            validate_matlab_metric_table_reports(metrics_dir)
        )
        manifest["classical_statistics_input"] = {
            "noddi_compatibility_table": {
                "name": classical_file.name,
                "size_bytes": classical_file.stat().st_size,
                "sha256": sha256_file(classical_file),
            },
            "active_seven_metric_table": {
                "name": combined_classical_file.name,
                "size_bytes": combined_classical_file.stat().st_size,
                "sha256": sha256_file(combined_classical_file),
            },
        }
    output_dir = paths["stats"]
    if not args.dry_run:
        output_dir.mkdir(parents=True, exist_ok=True)
    expression = (
        f"addpath('{matlab_literal(MATLAB_RUNNER.parent)}','-begin'); "
        f"run_all_qmri_manuscript_models('statsDir','{matlab_literal(stats_dir)}',"
        f"'metricsDir','{matlab_literal(metrics_dir)}',"
        f"'classicalFile','{matlab_literal(classical_file)}',"
        f"'outputDir','{matlab_literal(output_dir)}');"
    )
    run_command(
        [str(paths["matlab"]), "-batch", expression],
        manifest=manifest,
        stage="manuscript_aligned_statistics",
        dry_run=args.dry_run,
    )
    models_csv = output_dir / "all_qmri_manuscript_models.csv"
    noddi_compatibility_csv = output_dir / "noddi_manuscript_aligned_models.csv"
    full_models_csv = output_dir / "figure6_full_models.csv"
    matched_models_csv = output_dir / "figure6_matched_models.csv"
    figure_dir = output_dir / "reviewer_results" / "figures"
    table_dir = output_dir / "reviewer_results" / "tables"
    if not args.dry_run:
        manifest["active_model_output_validation"] = validate_model_outputs(
            models_csv,
            noddi_compatibility_csv,
            full_models_csv,
            matched_models_csv,
        )
        analysis_report_path = output_dir / "all_qmri_manuscript_analysis_report.json"
        resolve_existing(analysis_report_path, "all-qMRI MATLAB analysis report")
        analysis_report = json.loads(analysis_report_path.read_text(encoding="utf-8"))
        reference_validation = analysis_report.get("unified_reference_validation", {})
        if (
            analysis_report.get("analysis_scope") != "all-metrics"
            or int(analysis_report.get("model_rows", 0)) != 196
            or int(analysis_report.get("global_family_size", 0)) != 98
            or reference_validation.get("status") != "PASS"
            or int(reference_validation.get("legacy_adjusted_rows_checked", 0)) != 56
            or int(reference_validation.get("noddi_rows_checked", 0)) != 84
            or analysis_report.get("classical_legacy_validation", {}).get("status")
            != "PASS"
        ):
            raise ReproductionError(
                "The all-qMRI MATLAB analysis report failed its scientific validation gate"
            )
        bundled_models = (
            PACKAGE_ROOT
            / "results"
            / "model_results"
            / "all_qmri_manuscript_models.csv"
        )
        generated_models_hash = sha256_file(models_csv)
        bundled_models_hash = sha256_file(bundled_models)
        manifest["active_model_output_validation"]["bundled_result_comparison"] = {
            "comparison": (
                "Generated and bundled 196-row scientific aggregates; hashes are retained "
                "as provenance and MATLAB numeric/text validation is authoritative"
            ),
            "generated_sha256": generated_models_hash,
            "bundled_sha256": bundled_models_hash,
            "byte_identical": generated_models_hash == bundled_models_hash,
            "matlab_validation": reference_validation,
            "status": "PASS",
        }
    table_expression = (
        f"addpath('{matlab_literal(MATLAB_TABLE_RUNNER.parent)}','-begin'); "
        f"generate_manuscript_tables("
        f"'StatsDir','{matlab_literal(stats_dir)}',"
        f"'ModelFile','{matlab_literal(models_csv)}',"
        f"'LesionModelFile','{matlab_literal(PACKAGE_ROOT / 'results' / 'model_results' / 'lesionload_manuscript_models.csv')}',"
        f"'OutputDir','{matlab_literal(table_dir)}');"
    )
    run_command(
        [str(paths["matlab"]), "-batch", table_expression],
        manifest=manifest,
        stage="plain_manuscript_tables",
        dry_run=args.dry_run,
    )
    if not args.dry_run:
        manifest["manuscript_table_validation"] = validate_manuscript_table_outputs(
            table_dir
        )
    figure_expression = (
        f"addpath('{matlab_literal(MATLAB_FIGURE_RUNNER.parent)}','-begin'); "
        f"figure6_r2_comparison_sdmt_msfc_noddi("
        f"'fullFile','{matlab_literal(full_models_csv)}',"
        f"'sameNFile','{matlab_literal(matched_models_csv)}',"
        f"'outputDir','{matlab_literal(figure_dir)}',"
        f"'previewDir','{matlab_literal(figure_dir)}');"
    )
    run_command(
        [str(paths["matlab"]), "-batch", figure_expression],
        manifest=manifest,
        stage="manuscript_style_figures",
        dry_run=args.dry_run,
    )
    if not args.dry_run:
        figure_names = (
            "Figure6_with_SDMT_MSFC_NODDI",
            "Supplementary_Figure_S2_with_SDMT_MSFC_NODDI",
        )
        generated_pdfs = [figure_dir / f"{name}.pdf" for name in figure_names]
        generated_pngs = [figure_dir / f"{name}.png" for name in figure_names]
        output_records: dict[str, Any] = {}
        for generated_png in generated_pngs:
            resolve_existing(generated_png, "generated manuscript-style PNG")
            if (
                generated_png.stat().st_size < 1000
                or generated_png.read_bytes()[:8] != b"\x89PNG\r\n\x1a\n"
            ):
                raise ReproductionError(f"Invalid PNG output: {generated_png}")
            output_records[generated_png.name] = {
                "comparison": "active all-metric rendering; valid PNG and content hash recorded",
                "size_bytes": generated_png.stat().st_size,
                "sha256": sha256_file(generated_png),
            }
        for generated_pdf in generated_pdfs:
            resolve_existing(generated_pdf, "generated manuscript-style PDF")
            if (
                generated_pdf.stat().st_size < 1000
                or generated_pdf.read_bytes()[:5] != b"%PDF-"
            ):
                raise ReproductionError(f"Invalid PDF output: {generated_pdf}")
            output_records[generated_pdf.name] = {
                "comparison": "active all-metric rendering; valid PDF and content hash recorded",
                "size_bytes": generated_pdf.stat().st_size,
                "sha256": sha256_file(generated_pdf),
            }
        manifest["reviewer_output_validation"] = {
            "status": "PASS",
            "files": output_records,
            "note": (
                "Figures were generated from the active full-precision 196-row refit; "
                "legacy pixel identity is intentionally not required."
            ),
        }


def final_output_checks(paths: dict[str, Path], stage: str) -> dict[str, Any]:
    expected: list[Path] = []
    if stage in {"fit", "all"}:
        expected.append(paths["fit"] / "logs" / "run_status.csv")
    if stage in {"extract", "all"}:
        expected.extend(
            paths["tract"] / "results" / filename
            for filename in (
                "GroupTractT1_All.xlsx",
                "GroupTractT1_All.csv",
                "GroupTractMTR_All.xlsx",
                "GroupTractMTR_All.csv",
                "GroupTractFA_All.xlsx",
                "GroupTractFA_All.csv",
                "GroupTractMD_All.xlsx",
                "GroupTractMD_All.csv",
                "GroupTractNDI_All.csv",
                "GroupTractODI_All.csv",
                "GroupTractFWF_All.csv",
                "ClassicalRegionNODDI_All.csv",
                "ClassicalRegionAllMetrics.csv",
                "clinicalScore.xlsx",
                "icometrixLesionStats.xlsx",
            )
        )
    if stage in {"stats", "all"}:
        expected.extend(
            [
                paths["stats"] / "all_qmri_manuscript_models.csv",
                paths["stats"] / "noddi_manuscript_aligned_models.csv",
                paths["stats"] / "figure6_full_models.csv",
                paths["stats"] / "figure6_matched_models.csv",
                paths["stats"] / "fa_md_equivalence_validation.json",
                paths["stats"] / "all_qmri_manuscript_analysis_report.json",
                paths["stats"] / "reviewer_results" / "figures" / "Figure6_with_SDMT_MSFC_NODDI.pdf",
                paths["stats"] / "reviewer_results" / "figures" / "Figure6_with_SDMT_MSFC_NODDI.png",
                paths["stats"] / "reviewer_results" / "figures" / "Supplementary_Figure_S2_with_SDMT_MSFC_NODDI.pdf",
                paths["stats"] / "reviewer_results" / "figures" / "Supplementary_Figure_S2_with_SDMT_MSFC_NODDI.png",
                paths["stats"] / "reviewer_results" / "tables" / "Main_Table_1.csv",
                paths["stats"] / "reviewer_results" / "tables" / "Main_Table_2.csv",
                paths["stats"] / "reviewer_results" / "tables" / "Supplementary_Table_3.csv",
                paths["stats"] / "reviewer_results" / "tables" / "Supplementary_Table_4.csv",
                paths["stats"] / "reviewer_results" / "tables" / "Supplementary_Table_5.csv",
            ]
        )
    missing = [path for path in expected if not path.is_file()]
    if missing:
        raise ReproductionError(
            "Expected outputs are missing:\n  " + "\n  ".join(str(path) for path in missing)
        )
    existing = expected
    return {
        "expected_core_files": len(expected),
        "existing_files": len(existing),
        "status": "PASS",
        "outputs": [
            {
                "path": str(path),
                "size_bytes": path.stat().st_size,
                "sha256": sha256_file(path),
            }
            for path in existing
        ],
    }


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "stage",
        nargs="?",
        choices=("check", "fit", "extract", "stats", "all"),
        default="check",
        help="Pipeline stage to run (default: check).",
    )
    parser.add_argument("--bids-root", type=Path, help="Private MsBIDS root (read-only).")
    parser.add_argument("--work-root", type=Path, help="Local/private output root.")
    parser.add_argument("--python", type=Path, help="Python executable with helper/config/requirements.txt.")
    parser.add_argument("--ants", type=Path, help="antsApplyTransforms executable.")
    parser.add_argument("--matlab", type=Path, help="MATLAB executable.")
    parser.add_argument(
        "--freesurfer-matlab",
        type=Path,
        help="FreeSurfer MATLAB utilities directory containing MRIread.m.",
    )
    parser.add_argument(
        "--group-mask-overlay",
        type=Path,
        help=(
            "Read-only sub-*/groupTract mask overlay used only for masks absent "
            "from --bids-root (extract/all stages)."
        ),
    )
    parser.add_argument(
        "--stats-dir",
        type=Path,
        help=(
            "For a standalone stats run, local directory containing clinicalScore, "
            "and icometrixLesionStats. Defaults to <work-root>/tract/results."
        ),
    )
    parser.add_argument(
        "--metrics-dir",
        type=Path,
        help=(
            "For a standalone stats run, local directory containing the actively "
            "generated GroupTract T1/MTR/FA/MD workbooks, NDI/ODI/FWF CSVs, "
            "and classical-region tables. Defaults to <work-root>/tract/results."
        ),
    )
    parser.add_argument(
        "--classical-file",
        type=Path,
        help=(
            "Explicit ClassicalRegionNODDI_All.csv for the stats stage; "
            "defaults to the metrics directory."
        ),
    )
    parser.add_argument("--workers", type=int, default=4, choices=range(1, 5))
    parser.add_argument("--snapshot-workers", type=int, default=3)
    parser.add_argument("--qc-figures", action="store_true")
    parser.add_argument("--resume", action="store_true", help="Allow a non-empty work root and validate reuse.")
    parser.add_argument(
        "--force",
        action="store_true",
        help="Regenerate local map transforms and MATLAB aggregate tables.",
    )
    parser.add_argument("--dry-run", action="store_true", help="Print stage commands without running them.")
    parser.add_argument(
        "--package-only",
        action="store_true",
        help="For check: validate code/dependencies without requiring the private BIDS tree.",
    )
    return parser


def main(argv: Sequence[str] | None = None) -> int:
    parser = build_parser()
    args = parser.parse_args(argv)
    if args.snapshot_workers < 1:
        parser.error("--snapshot-workers must be positive")
    if args.package_only and args.stage != "check":
        parser.error("--package-only is valid only with the check stage")
    if (
        args.stats_dir is not None
        or args.metrics_dir is not None
        or args.classical_file is not None
    ) and args.stage != "stats":
        parser.error(
            "--stats-dir, --metrics-dir, and --classical-file overrides are valid "
            "only with the stats stage"
        )
    if args.freesurfer_matlab is not None and args.stage not in {"check", "extract", "all"}:
        parser.error(
            "--freesurfer-matlab is valid only with the check, extract, or all stage"
        )
    if args.group_mask_overlay is not None and args.stage not in {"extract", "all"}:
        parser.error("--group-mask-overlay is valid only with extract or all")

    manifest: dict[str, Any] = {
        "schema_version": 1,
        "started_utc": utc_now(),
        "stage_requested": args.stage,
        "package_root": str(PACKAGE_ROOT),
        "package_hashes": package_hashes(),
    }
    work_root: Path | None = None
    try:
        bids_root = None
        if args.bids_root is not None:
            bids_root = resolve_existing(args.bids_root, "BIDS root", directory=True)
        group_mask_overlay = None
        if args.group_mask_overlay is not None:
            group_mask_overlay = resolve_existing(
                args.group_mask_overlay, "grouped-mask overlay", directory=True
            )
        if args.stage in {"fit", "extract", "all"} and bids_root is None:
            raise ReproductionError("--bids-root is required for this stage")

        if args.stage != "check":
            if args.work_root is None:
                raise ReproductionError("--work-root is required for this stage")
            work_root = validate_work_root(args.work_root, bids_root)
            if group_mask_overlay is not None and (
                work_root == group_mask_overlay
                or is_relative_to(work_root, group_mask_overlay)
                or is_relative_to(group_mask_overlay, work_root)
            ):
                raise ReproductionError(
                    "--work-root and --group-mask-overlay must be separate trees"
                )
            if (
                args.stage in {"fit", "all"}
                and work_root.exists()
                and any(work_root.iterdir())
                and not args.resume
            ):
                raise ReproductionError(
                    f"Work root is not empty: {work_root}. Use a new directory or pass --resume."
                )
            if not args.dry_run:
                work_root.mkdir(parents=True, exist_ok=True)

        python = locate_executable(args.python, (sys.executable,), "Python")
        ants = None
        if args.stage in {"check", "extract", "all"}:
            try:
                ants = locate_executable(
                    args.ants,
                    ("antsApplyTransforms", "/opt/ANTs/bin/antsApplyTransforms"),
                    "antsApplyTransforms",
                )
            except ReproductionError:
                if args.stage != "check" or not args.package_only:
                    raise
        matlab = None
        if args.stage in {"check", "extract", "stats", "all"}:
            try:
                matlab = locate_executable(
                    args.matlab,
                    (
                        "matlab",
                        "/Applications/MATLAB_R2024b.app/bin/matlab",
                    ),
                    "MATLAB",
                )
            except ReproductionError:
                if args.stage != "check" or not args.package_only:
                    raise
        freesurfer_matlab = None
        if args.stage in {"check", "extract", "all"}:
            try:
                freesurfer_matlab = locate_freesurfer_matlab(args.freesurfer_matlab)
            except ReproductionError:
                if (
                    args.freesurfer_matlab is not None
                    or args.stage != "check"
                    or not args.package_only
                ):
                    raise

        report = dependency_report(python)
        manifest["environment"] = report
        if report["mismatches"]:
            raise ReproductionError(
                "Python environment does not match helper/config/requirements.txt: "
                + json.dumps(report["mismatches"], sort_keys=True)
            )
        manifest["multiplicity_validation"] = verify_multiplicity_accounting()
        if bids_root is not None and not args.package_only:
            manifest["source_validation"] = validate_source_tree(bids_root)
            if args.stage in {"check", "extract", "all"}:
                manifest["upstream_statistics_validation"] = validate_aggregate_inputs(
                    bids_root / "derivatives" / "derivativesStats"
                )

        if args.stage == "check":
            manifest.update(
                status="PASS",
                ants=str(ants) if ants else "NOT CHECKED",
                matlab=str(matlab) if matlab else "NOT CHECKED",
                freesurfer_matlab=(
                    str(freesurfer_matlab) if freesurfer_matlab else "NOT CHECKED"
                ),
                finished_utc=utc_now(),
            )
            print(json.dumps(manifest, indent=2, sort_keys=True))
            return 0

        assert work_root is not None
        paths = {
            "bids": bids_root,
            "work": work_root,
            "fit": work_root / "fit",
            "tract": work_root / "tract",
            "stats": work_root / "stats",
            "python": python,
            "ants": ants,
            "matlab": matlab,
            "freesurfer_matlab": freesurfer_matlab,
            "group_mask_overlay": group_mask_overlay,
        }
        if args.stage in {"fit", "all"}:
            run_fit(args, paths, manifest)
        if args.stage in {"extract", "all"}:
            if not paths["fit"].is_dir() and not args.dry_run:
                raise ReproductionError(f"NODDI fit root is missing: {paths['fit']}")
            run_extract(args, paths, manifest)
        if args.stage in {"stats", "all"}:
            run_stats(args, paths, manifest)
        if not args.dry_run:
            manifest["output_checks"] = final_output_checks(paths, args.stage)
        manifest.update(status="PASS", finished_utc=utc_now())
        if args.dry_run:
            print("\nPASS. Dry run completed; no files were written by the public runner.")
        else:
            atomic_json(work_root / "run_manifest.json", manifest)
            print(f"\nPASS. Private run manifest: {work_root / 'run_manifest.json'}")
        return 0
    except Exception as exc:
        manifest.update(
            status="FAIL",
            finished_utc=utc_now(),
            error=f"{type(exc).__name__}: {exc}",
        )
        if work_root is not None and not args.dry_run:
            atomic_json(work_root / "run_manifest.json", manifest)
        print(f"ERROR: {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
