#!/usr/bin/env python3
"""Stage read-only source images locally; build Motor masks and extract seven metrics."""
from __future__ import annotations
import argparse, csv, json, os, shutil
from concurrent.futures import ProcessPoolExecutor
from collections import Counter
from pathlib import Path
import nibabel as nib
import numpy as np
from scipy.ndimage import convolve
from openpyxl import load_workbook

METRICS=("T1","MTR","FA","MD","NDI","ODI","FWF")
COLS=("MotorL_Tail","MotorR_Tail","MotorL_NAWM","MotorR_NAWM")
TAIL={"T1":90,"MTR":10,"FA":10,"MD":90,"NDI":10,"ODI":90,"FWF":90}

def subjects_from_clinical(path):
    ws=load_workbook(path,read_only=True,data_only=True).active
    rows=ws.iter_rows(values_only=True); h=next(rows); index=h.index("SubjectID")
    ids=[str(r[index]) for r in rows if r[index]]
    if len(ids)!=132 or len(set(ids))!=132: raise ValueError("Expected 132 unique clinical subjects")
    return ids

def subjects_from_metric_table(path):
    ws=load_workbook(path,read_only=True,data_only=True).active
    rows=ws.iter_rows(values_only=True); h=next(rows); index=h.index('SubjectID')
    ids=[str(r[index]) for r in rows if r[index]]
    if len(ids)!=132 or len(set(ids))!=132: raise ValueError('Expected 132 unique manuscript imaging subjects')
    return ids

def copy_new(src,dst):
    if not src.is_file(): return False
    if dst.exists():
        if dst.stat().st_size != src.stat().st_size: raise FileExistsError(dst)
        return True
    dst.parent.mkdir(parents=True,exist_ok=True)
    with src.open('rb') as fi, dst.open('xb') as fo: shutil.copyfileobj(fi,fo,1024*1024)
    if dst.stat().st_size!=src.stat().st_size: raise IOError(f"Partial copy: {dst}")
    return True

def stage_one(source,stage,subject,tracts):
    d=stage/subject; root=source/'derivatives'
    for tract in tracts:
        copy_new(root/'TractoFlow_post'/subject/'trkTract'/f'{tract}.nii.gz',d/'trkTract'/f'{tract}.nii.gz')
    for hemi in ('lh','rh'):
        copy_new(root/'freesurfer'/subject/'mri'/f'{hemi}.brainMask.nii.gz',d/f'{hemi}.brainMask.nii.gz')
    for name in (f'{subject}_ses-01_space-individual_T1map.nii.gz',f'{subject}_ses-01_space-individual_MTRmap.nii.gz'):
        copy_new(root/'maps'/subject/name,d/name)

def mask_image(path):
    im=nib.load(str(path)); return np.asarray(im.dataobj)>0,im

def orient(mask): return np.flip(np.flip(np.transpose(mask,(0,2,1)),1),2)

def oriented_affine(im):
    transform=np.array([[1,0,0,0],[0,0,-1,im.shape[1]-1],[0,-1,0,im.shape[2]-1],[0,0,0,1]],float)
    return im.affine@transform

def require_geometry(im,ref,label,tol=1e-4):
    if im.shape!=ref.shape or np.max(np.abs(im.affine-ref.affine))>tol: raise ValueError(f'{label} geometry mismatch')

def save_mask(mask,ref,path):
    if path.exists(): raise FileExistsError(path)
    header=ref.header.copy(); header.set_data_dtype(np.uint8)
    out=nib.Nifti1Image(mask.astype(np.uint8),ref.affine,header)
    out.set_qform(ref.get_qform(),int(ref.header['qform_code'])); out.set_sform(ref.get_sform(),int(ref.header['sform_code']))
    path.parent.mkdir(parents=True,exist_ok=True); nib.save(out,str(path))

def values(data,roi,metric,summary):
    v=np.asarray(data[roi]); noddi=metric in ('NDI','ODI','FWF')
    if noddi and not np.all(np.isfinite(v)): raise ValueError('NODDI ROI nonfinite')
    v=v[np.isfinite(v)]
    if not v.size: return float('nan') if noddi else 0.0
    if summary=='Tail': return float(np.percentile(v,TAIL[metric],method='hazen'))
    return float(np.mean(v,dtype=np.float64))

def extract_one(cfg,stage,subject,tracts,mask_dir):
    local=Path(cfg['local_snapshot']); noddi=Path(cfg['noddi_maps_dir']); sd=stage/subject
    ref=nib.load(str(local/'derivatives/maps'/subject/f'{subject}_space-individual_FA.nii.gz'))
    combined=None; base=None; present=[]; empty=[]; missing=[]
    for tract in tracts:
        path=sd/'trkTract'/f'{tract}.nii.gz'
        if not path.is_file(): missing.append(tract); continue
        m,im=mask_image(path)
        if base is None: base=im; combined=np.zeros(im.shape,bool)
        else: require_geometry(im,base,tract)
        present.append(tract)
        if not m.any(): empty.append(tract)
        combined|=m
    if base is None or not combined.any():
        missing_values={metric:{column:float('nan') for column in COLS} for metric in METRICS}
        return missing_values,{'present':present,'empty':empty,'missing':missing,
            'combined_voxels':0,'left_voxels':0,'right_voxels':0,
            'left_centroid_x':None,'right_centroid_x':None,'split_exact':None,
            'geometry_pass':None,'invalid_metric_maps':{},
            'status':'UNAVAILABLE_NO_NONEMPTY_COMPONENT'}
    lh,lh_im=mask_image(sd/'lh.brainMask.nii.gz'); rh,rh_im=mask_image(sd/'rh.brainMask.nii.gz')
    require_geometry(rh_im,lh_im,'hemisphere masks')
    lh=orient(lh); rh=orient(rh)
    if combined.shape!=lh.shape or np.max(np.abs(base.affine-oriented_affine(lh_im)))>1e-4: raise ValueError('Hemisphere versus tract geometry')
    if orient(combined).shape!=ref.shape or np.max(np.abs(oriented_affine(base)-ref.affine))>1e-4: raise ValueError('Motor versus map geometry')
    kernel=np.ones((3,3,3),np.uint8)
    lc=convolve(lh.astype(np.uint8),kernel,mode='constant'); rc=convolve(rh.astype(np.uint8),kernel,mode='constant')
    left_only=lh & ~rh; right_only=rh & ~lh; ambiguous=~(left_only|right_only)
    assign_left=left_only | (ambiguous & (lc>=rc))
    left=combined & assign_left; right=combined & ~assign_left
    if not left.any() or not right.any() or np.any(left&right) or not np.array_equal(left|right,combined): raise ValueError('Invalid Motor L/R split')
    for name,m in [('Motor',combined),('MotorL',left),('MotorR',right)]: save_mask(m,base,mask_dir/f'{name}.nii.gz')
    oriented={h:orient(m) for h,m in [('L',left),('R',right)]}
    for h,m in oriented.items():
        xyz=np.argwhere(m); world=nib.affines.apply_affine(ref.affine,xyz)
        if h=='L': left_x=float(world[:,0].mean())
        else: right_x=float(world[:,0].mean())
    if not left_x<right_x: raise ValueError('Motor anatomical L/R centroid reversal')
    lesion_path=local/'derivatives/lesionMask'/subject/'ses-01'/f'{subject}_ses-01_desc-lesionManual_mask.nii.gz'
    if subject.startswith('sub-C'): lesion=np.zeros(ref.shape,bool)
    elif lesion_path.is_file():
        lesion_im=nib.load(str(lesion_path)); require_geometry(lesion_im,ref,'lesion'); lesion=np.asarray(lesion_im.dataobj)>0
    else: lesion=np.zeros(ref.shape,bool)
    support_path=noddi/subject/f'{subject}_space-individual_NODDI-support.nii.gz'
    support_im=nib.load(str(support_path)); require_geometry(support_im,ref,'NODDI support'); support=np.asarray(support_im.dataobj)>0
    result={}; invalid={}
    for metric in METRICS:
        if metric in ('T1','MTR'): p=sd/f'{subject}_ses-01_space-individual_{metric}map.nii.gz'
        elif metric in ('FA','MD'): p=local/'derivatives/maps'/subject/f'{subject}_space-individual_{metric}.nii.gz'
        else: p=noddi/subject/f'{subject}_space-individual_{metric}.nii.gz'
        if not p.is_file(): result[metric]={c:float('nan') for c in COLS}; invalid[metric]='missing_map'; continue
        im=nib.load(str(p)); require_geometry(im,ref,metric); data=np.asarray(im.dataobj)
        if metric in ('NDI','ODI','FWF'):
            vals=data[support]
            if not np.all(np.isfinite(vals)) or np.any(vals < -1e-6) or np.any(vals > 1+1e-6): raise ValueError(f'{metric} invalid fitted support')
        mrow={}
        for h,m in oriented.items():
            roi=m & (support if metric in ('NDI','ODI','FWF') else True)
            if not np.any(roi): mrow[f'Motor{h}_Tail']=float('nan'); mrow[f'Motor{h}_NAWM']=float('nan'); continue
            mrow[f'Motor{h}_Tail']=values(data,roi,metric,'Tail')
            nawm=roi if subject.startswith('sub-C') else roi & ~lesion
            mrow[f'Motor{h}_NAWM']=values(data,nawm,metric,'NAWM')
            if metric in ('NDI','ODI','FWF') and any(not (0<=mrow[f'Motor{h}_{s}']<=1) for s in ('Tail','NAWM')): raise ValueError(f'{metric} aggregate out of bounds')
        result[metric]=mrow
    qc={'present':present,'empty':empty,'missing':missing,'combined_voxels':int(combined.sum()),'left_voxels':int(left.sum()),'right_voxels':int(right.sum()),'left_centroid_x':left_x,'right_centroid_x':right_x,'split_exact':True,'geometry_pass':True,'invalid_metric_maps':invalid,'status':'PASS'}
    return result,qc

def extract_task(args):
    return extract_one(*args)

def main():
    p=argparse.ArgumentParser(); p.add_argument('--config',type=Path,required=True); p.add_argument('--stage-dir',type=Path,required=True); p.add_argument('--output-dir',type=Path,required=True); p.add_argument('--stage-only',action='store_true'); p.add_argument('--workers',type=int,default=1); a=p.parse_args()
    cfg=json.loads(a.config.read_text()); source=Path(cfg['source_bids_root']).resolve(); stage=a.stage_dir.resolve(); out=a.output_dir.resolve()
    scratch=Path(cfg['scratch_root']).resolve()
    if not stage.is_relative_to(scratch) or not out.is_relative_to(scratch) or source==stage or stage.is_relative_to(source): raise ValueError('Unsafe paths')
    if out.exists(): raise FileExistsError(out)
    tracts=[x.strip() for x in (Path(__file__).resolve().parents[1]/'config/motor.txt').read_text().splitlines() if x.strip() and not x.startswith(('#','['))]
    if len(tracts)!=13 or len(set(tracts))!=13: raise ValueError('Motor definition must contain 13 unique tracts')
    subjects=subjects_from_metric_table(Path(cfg['metrics_dir'])/'GroupTractT1_All.xlsx')
    stage.mkdir(parents=True,exist_ok=True)
    if not (stage/'STAGED_IMAGING.json').is_file():
        for i,subject in enumerate(subjects,1):
            stage_one(source,stage,subject,tracts)
            if i%10==0: print(f'Staged {i}/{len(subjects)}',flush=True)
        with (stage/'STAGED_IMAGING.json').open('x') as f: json.dump({'subjects':len(subjects),'components':len(tracts),'cohort':'frozen manuscript imaging tables'},f)
    if a.stage_only: return
    out.mkdir(parents=True)
    rows={m:[] for m in METRICS}; qc=[]
    tasks=[(cfg,stage,subject,tracts,out/'masks'/subject) for subject in subjects]
    if a.workers<1 or a.workers>4: raise ValueError('workers must be between 1 and 4')
    if a.workers==1:
        extracted=map(extract_task,tasks)
    else:
        pool=ProcessPoolExecutor(max_workers=a.workers)
        extracted=pool.map(extract_task,tasks)
    for i,(subject,(metric_values,record)) in enumerate(zip(subjects,extracted),1):
        qc.append({'subject':subject,**record})
        for m in METRICS: rows[m].append({'SubjectID':subject,**metric_values[m]})
        if i%10==0: print(f'Extracted {i}/{len(subjects)}',flush=True)
    if a.workers>1: pool.shutdown()
    (out/'metrics').mkdir()
    for m in METRICS:
        with (out/'metrics'/f'Motor{m}_All.csv').open('x',newline='') as f:
            w=csv.DictWriter(f,fieldnames=['SubjectID',*COLS]);w.writeheader();w.writerows(rows[m])
    with (out/'subject_qc.json').open('x') as f:json.dump(qc,f,indent=2)
    print(f'Extraction complete: {len(subjects)} subjects',flush=True)
if __name__=='__main__': main()
